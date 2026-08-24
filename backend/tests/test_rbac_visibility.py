"""End-to-end visibility matrix for the distributor → reseller hierarchy.

Everything here goes over real HTTP against the real app: real routes, real
dependencies, real SQL. Unit tests can prove the scope resolver returns the
right list of ids; only this can prove that every read path actually uses it,
which is the part that breaks when someone adds an endpoint.

The world each test builds:

    Distributor  ──┬── Reseller A   (partner, parent = Distributor)
                   └── Reseller B   (partner, parent = Distributor)
    Independent      (partner, no parent)

and the three questions the feature has to answer:

    - the distributor sees A's and B's pipeline           (it is above them)
    - A does not see B's, and B does not see A's          (they are siblings)
    - neither sees the distributor's                      (it is above them)

plus the two boundaries that keep the widening honest: reads went
company-wide but writes did not, and money data (deals) did not follow the
tree even though pipeline data did.
"""
import pytest

from app.models.company import CompanyType
from app.models.user import UserRole
from tests.conftest import (
    auth_header,
    make_company,
    make_opportunity,
    make_user,
    requires_db,
)

pytestmark = [requires_db, pytest.mark.asyncio]


class World:
    """The fixture world, with the handful of ids each test asserts on."""


async def build_world(db):
    """Distributor with two resellers, plus an unrelated independent partner.

    Every company gets its own opportunity submitted by its own user, so
    "who can see this row" has exactly one right answer per row.
    """
    w = World()
    w.admin = await make_user(db, role=UserRole.ADMIN, is_superadmin=True)

    w.distributor = await make_company(
        db, channel_manager_id=w.admin.id, company_type=CompanyType.DISTRIBUTOR
    )
    w.reseller_a = await make_company(
        db, channel_manager_id=w.admin.id, company_type=CompanyType.PARTNER,
        parent_distributor_id=w.distributor.id,
    )
    w.reseller_b = await make_company(
        db, channel_manager_id=w.admin.id, company_type=CompanyType.PARTNER,
        parent_distributor_id=w.distributor.id,
    )
    w.independent = await make_company(
        db, channel_manager_id=w.admin.id, company_type=CompanyType.PARTNER
    )

    w.dist_user = await make_user(
        db, role=UserRole.PARTNER, company_id=w.distributor.id
    )
    w.a_user = await make_user(db, role=UserRole.PARTNER, company_id=w.reseller_a.id)
    # A colleague inside reseller A, for the company-wide-reads assertion.
    w.a_colleague = await make_user(
        db, role=UserRole.PARTNER, company_id=w.reseller_a.id
    )
    w.b_user = await make_user(db, role=UserRole.PARTNER, company_id=w.reseller_b.id)
    w.indep_user = await make_user(
        db, role=UserRole.PARTNER, company_id=w.independent.id
    )

    w.dist_opp = await make_opportunity(
        db, company_id=w.distributor.id, submitted_by=w.dist_user.id
    )
    w.a_opp = await make_opportunity(
        db, company_id=w.reseller_a.id, submitted_by=w.a_user.id
    )
    w.b_opp = await make_opportunity(
        db, company_id=w.reseller_b.id, submitted_by=w.b_user.id
    )
    w.indep_opp = await make_opportunity(
        db, company_id=w.independent.id, submitted_by=w.indep_user.id
    )

    # The app runs on its own session, so nothing above is visible to a
    # request until it is committed.
    await db.commit()
    return w


async def visible_opportunity_ids(client, headers) -> set[int]:
    r = await client.get("/api/v1/opportunities?page_size=100", headers=headers)
    assert r.status_code == 200, r.text[:300]
    return {item["id"] for item in r.json()["items"]}


# ---------------------------------------------------------------------------
# The list: who sees whose pipeline
# ---------------------------------------------------------------------------

class TestOpportunityListScope:
    async def test_distributor_sees_both_resellers(self, client, db):
        w = await build_world(db)
        seen = await visible_opportunity_ids(
            client, auth_header(w.dist_user)
        )
        assert {w.dist_opp.id, w.a_opp.id, w.b_opp.id} <= seen

    async def test_distributor_does_not_see_an_unrelated_company(self, client, db):
        w = await build_world(db)
        seen = await visible_opportunity_ids(
            client, auth_header(w.dist_user)
        )
        assert w.indep_opp.id not in seen

    async def test_resellers_do_not_see_each_other(self, client, db):
        # The requirement in one assertion: siblings are invisible to each
        # other in both directions.
        w = await build_world(db)
        a_seen = await visible_opportunity_ids(
            client, auth_header(w.a_user)
        )
        b_seen = await visible_opportunity_ids(
            client, auth_header(w.b_user)
        )
        assert w.b_opp.id not in a_seen
        assert w.a_opp.id not in b_seen

    async def test_reseller_does_not_see_its_parent(self, client, db):
        # Visibility walks downward only.
        w = await build_world(db)
        a_seen = await visible_opportunity_ids(
            client, auth_header(w.a_user)
        )
        assert w.dist_opp.id not in a_seen
        assert w.a_opp.id in a_seen

    async def test_company_id_filter_cannot_escape_the_scope(self, client, db):
        # The filter narrows within the scope; it is not a way out of it.
        w = await build_world(db)
        headers = auth_header(w.a_user)
        r = await client.get(
            f"/api/v1/opportunities?company_id={w.reseller_b.id}", headers=headers
        )
        assert r.status_code == 200
        assert r.json()["items"] == []


# ---------------------------------------------------------------------------
# The detail route: the same scope, per record
# ---------------------------------------------------------------------------

class TestOpportunityDetailScope:
    async def test_distributor_can_open_a_reseller_opportunity(self, client, db):
        w = await build_world(db)
        headers = auth_header(w.dist_user)
        r = await client.get(f"/api/v1/opportunities/{w.a_opp.id}", headers=headers)
        assert r.status_code == 200

    async def test_reseller_cannot_open_a_siblings_opportunity(self, client, db):
        w = await build_world(db)
        headers = auth_header(w.a_user)
        r = await client.get(f"/api/v1/opportunities/{w.b_opp.id}", headers=headers)
        assert r.status_code == 403

    async def test_reseller_cannot_open_its_parents_opportunity(self, client, db):
        w = await build_world(db)
        headers = auth_header(w.a_user)
        r = await client.get(f"/api/v1/opportunities/{w.dist_opp.id}", headers=headers)
        assert r.status_code == 403


# ---------------------------------------------------------------------------
# Reads widened, writes did not
# ---------------------------------------------------------------------------

class TestReadsWidenedWritesDidNot:
    async def test_colleague_can_read_a_teammates_opportunity(self, client, db):
        w = await build_world(db)
        headers = auth_header(w.a_colleague)
        r = await client.get(f"/api/v1/opportunities/{w.a_opp.id}", headers=headers)
        assert r.status_code == 200

    async def test_colleague_cannot_edit_a_teammates_opportunity(self, client, db):
        # Reads went company-wide; editing is still the submitter's alone.
        w = await build_world(db)
        headers = auth_header(w.a_colleague)
        r = await client.put(
            f"/api/v1/opportunities/{w.a_opp.id}",
            headers=headers,
            json={"name": "Renamed by a colleague"},
        )
        assert r.status_code == 403

    async def test_distributor_cannot_edit_a_reseller_opportunity(self, client, db):
        w = await build_world(db)
        headers = auth_header(w.dist_user)
        r = await client.put(
            f"/api/v1/opportunities/{w.a_opp.id}",
            headers=headers,
            json={"name": "Renamed by the distributor"},
        )
        assert r.status_code == 403


# ---------------------------------------------------------------------------
# Money data does not follow the tree
# ---------------------------------------------------------------------------

class TestDealsStayInsideTheCompany:
    async def test_distributor_does_not_see_reseller_deals(self, client, db):
        # Deal registration carries exclusivity and commission. A distributor
        # reads its resellers' pipeline and nothing else, so a deal registered
        # by a reseller must not appear.
        w = await build_world(db)

        a_headers = auth_header(w.a_user)
        created = await client.post(
            "/api/v1/dashboard/deals",
            headers=a_headers,
            json={
                "customer_name": "Sibling-Safe Customer",
                "deal_description": "Reseller A's own deal",
                "estimated_value": 120000,
                "expected_close_date": "2027-09-30",
            },
        )
        assert created.status_code == 201, created.text[:300]
        deal_id = created.json()["id"]

        dist_headers = auth_header(w.dist_user)
        r = await client.get("/api/v1/dashboard/deals?page_size=100", headers=dist_headers)
        assert r.status_code == 200
        assert deal_id not in {d["id"] for d in r.json()["items"]}

    async def test_colleague_does_see_company_deals(self, client, db):
        # The other half of the same rule: company-wide, just not down the
        # tree. Without this the test above would pass on a scope that is
        # simply still per-user.
        w = await build_world(db)

        a_headers = auth_header(w.a_user)
        created = await client.post(
            "/api/v1/dashboard/deals",
            headers=a_headers,
            json={
                "customer_name": "Colleague-Visible Customer",
                "deal_description": "Reseller A's own deal",
                "estimated_value": 90000,
                "expected_close_date": "2027-10-31",
            },
        )
        assert created.status_code == 201, created.text[:300]
        deal_id = created.json()["id"]

        colleague_headers = auth_header(w.a_colleague)
        r = await client.get(
            "/api/v1/dashboard/deals?page_size=100", headers=colleague_headers
        )
        assert r.status_code == 200
        assert deal_id in {d["id"] for d in r.json()["items"]}


# ---------------------------------------------------------------------------
# Exports must not be a way around the list
# ---------------------------------------------------------------------------

class TestExportsMatchTheList:
    async def test_reseller_export_excludes_a_sibling(self, client, db):
        # An export that ignored the scope would hand over exactly the rows
        # the UI refuses to show.
        w = await build_world(db)
        headers = auth_header(w.a_user)
        r = await client.get("/api/v1/exports/opportunities.xlsx", headers=headers)
        assert r.status_code == 200

        from io import BytesIO

        from openpyxl import load_workbook

        ws = load_workbook(BytesIO(r.content)).active
        names = {
            str(row[1].value) for row in ws.iter_rows(min_row=2) if row[1].value
        }
        assert w.a_opp.name in names
        assert w.b_opp.name not in names
        assert w.dist_opp.name not in names


# ---------------------------------------------------------------------------
# Managing the hierarchy over the API
# ---------------------------------------------------------------------------

class TestHierarchyManagement:
    async def test_superadmin_can_link_a_partner_to_a_distributor(self, client, db):
        w = await build_world(db)
        headers = auth_header(w.admin)
        r = await client.put(
            f"/api/v1/companies/{w.independent.id}",
            headers=headers,
            json={"parent_distributor_id": w.distributor.id},
        )
        assert r.status_code == 200, r.text[:300]
        body = r.json()
        assert body["parent_distributor_id"] == w.distributor.id
        assert body["parent_distributor_name"] == w.distributor.name

    async def test_a_partner_cannot_be_parented_to_a_partner(self, client, db):
        w = await build_world(db)
        headers = auth_header(w.admin)
        r = await client.put(
            f"/api/v1/companies/{w.independent.id}",
            headers=headers,
            json={"parent_distributor_id": w.reseller_a.id},
        )
        assert r.status_code == 400

    async def test_a_distributor_cannot_be_given_a_parent(self, client, db):
        w = await build_world(db)
        headers = auth_header(w.admin)
        other_dist = await make_company(
            db, channel_manager_id=w.admin.id, company_type=CompanyType.DISTRIBUTOR
        )
        await db.commit()
        r = await client.put(
            f"/api/v1/companies/{w.distributor.id}",
            headers=headers,
            json={"parent_distributor_id": other_dist.id},
        )
        assert r.status_code == 400

    async def test_a_distributor_with_resellers_cannot_be_reclassified(self, client, db):
        w = await build_world(db)
        headers = auth_header(w.admin)
        r = await client.put(
            f"/api/v1/companies/{w.distributor.id}",
            headers=headers,
            json={"company_type": "partner"},
        )
        assert r.status_code == 409

    async def test_company_detail_lists_resellers_and_parent(self, client, db):
        w = await build_world(db)
        headers = auth_header(w.admin)

        parent_view = await client.get(
            f"/api/v1/companies/{w.distributor.id}", headers=headers
        )
        assert parent_view.status_code == 200
        reseller_ids = {r["id"] for r in parent_view.json()["resellers"]}
        assert {w.reseller_a.id, w.reseller_b.id} == reseller_ids

        child_view = await client.get(
            f"/api/v1/companies/{w.reseller_a.id}", headers=headers
        )
        assert child_view.status_code == 200
        assert child_view.json()["parent_distributor_id"] == w.distributor.id
        assert child_view.json()["resellers"] == []
